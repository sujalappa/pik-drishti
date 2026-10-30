"""Fill slides 3-6 of the Seva Innovation Challenge 2026 deck.

Run after build_slide2.py:  python build_slide2.py && python build_slides_3_6.py
Edits Pik-drishti_AI_Seva_2026.pptx in place.
"""

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
BUMP = 1  # pt added to every text size in this deck, for readability
DECK = HERE / "Pik-drishti_AI_Seva_2026.pptx"
DEMO_LINK = "pik-drishti.onrender.com"  # live demo on Render

INK = RGBColor(0x1F, 0x29, 0x37)
NAVY = RGBColor(0x1B, 0x2A, 0x5C)
ORANGE = RGBColor(0xF2, 0x8C, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x72, 0x80)
FONT = "Nirmala UI"


def fill(shape, paras, color=INK, size=10, bullets=False, space_after=3, bold=False, align=None):
    """Replace a text box's content, keeping its first paragraph/run formatting.
    paras: list of str or list of (text, bold) runs."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    proto_ppr = copy.deepcopy(first._p.pPr) if first._p.pPr is not None else None
    proto_rpr = first.runs[0]._r.rPr if first.runs else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        if i and proto_ppr is not None:
            p._p.insert(0, copy.deepcopy(proto_ppr))
        for text, b in ([(para, bold)] if isinstance(para, str) else para):
            r = p.add_run()
            if proto_rpr is not None:
                r._r.insert(0, copy.deepcopy(proto_rpr))
            r.text = text
            r.font.italic = False
            r.font.bold = b
            r.font.size = Pt(size + BUMP)
            r.font.color.rgb = color
        p.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        if bullets:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(Inches(0.13)))
            ppr.set("indent", str(-Inches(0.13)))
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for el in ppr.findall(qn(tag)):
                    ppr.remove(el)
            etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def textbox(slide, x, y, w, h, paras, size=9, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, b, col in ([(para, False, color)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = text
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size + BUMP), b, col
    return tb


def tick(slide, shape, size=11):
    """Fill a circle/checkbox orange and put a white tick on it."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    textbox(slide, shape.left / 914400, shape.top / 914400, shape.width / 914400, shape.height / 914400,
            [[("✔", True, WHITE)]], size=size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def remove(slide, names):
    for sh in list(slide.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def set_cell(cell, text, bold=False, color=INK, size=9):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    proto = p.runs[0]._r.rPr if p.runs else None
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    if proto is not None:
        r._r.insert(0, copy.deepcopy(proto))
    r.text = text
    r.font.italic, r.font.bold, r.font.size, r.font.color.rgb = False, bold, Pt(size + BUMP), color


prs = Presentation(DECK)
S = {i + 1: {sh.name: sh for sh in prs.slides[i].shapes} for i in range(len(prs.slides))}

# ---------------------------------------------------------------- slide 3: solution
s3, sh = prs.slides[2], S[3]
sh["Text 6"].width = Inches(2.3)
fill(sh["Text 6"], [[("Pik-drishti AI", True), (" turns free satellite data into village evidence for the ", False),
                     ("drought manual's triggers", True), (" and for crop-loss ", False), ("panchanama", True),
                     (".", False)]], size=9.5)
pic = s3.shapes.add_picture(str(HERE / "img" / "app_drought.png"), Inches(2.9), Inches(1.6), width=Inches(1.9))
pic.line.color.rgb = RGBColor(0xD9, 0xDE, 0xEA)
textbox(s3, 2.9, 1.6 + pic.height / 914400 + 0.03, 1.9, 0.18, ["Working prototype: drought map, Oct 2026"], size=6.5,
        color=MUTED)
fill(sh["Text 12"], ["A taluka, or village boundaries, plus a season or the date of a flood"], size=8.5,
     align=PP_ALIGN.CENTER)
fill(sh["Text 18"], ["Reads rainfall, radar and crop greenness; applies the drought manual's rules"], size=8.5,
     align=PP_ALIGN.CENTER)
fill(sh["Text 24"], ["Village map, ranked verification list, PDF, CSV and KML"], size=8.5, align=PP_ALIGN.CENTER)
table = sh["Table 0"].table
rows = [
    ("Cost", "Staff visit every village", "Free satellite data; small server"),
    ("Speed / time", "Weeks to 2 months a district", "Talukas ~1 min; villages ~2 min"),
    ("Reach / access", "Field staff miss farms", "All 48,926 villages, each season"),
    ("Ease of use", "Paper forms, manual totals", "Web map, PDF, KML; no GIS"),
]
for i, (label, old, new) in enumerate(rows, start=1):
    set_cell(table.cell(i, 0), label, bold=True, color=NAVY, size=9)
    set_cell(table.cell(i, 1), old, size=8.5)
    set_cell(table.cell(i, 2), new, size=8.5, bold=True)
fill(sh["Text 27"], [[("Follows the Government of India drought manual step by step and shows the evidence for "
                       "every village. Its Trigger-1 already matches the official 2026 list in ", False),
                      ("Jalgaon (15/15)", True), (" and ", False), ("Buldhana (13/13)", True), (".", False)]],
     size=9.5)

# ---------------------------------------------------------------- slide 4: feasibility
s4, sh = prs.slides[3], S[4]
for name in ("Shape 7", "Shape 9"):
    sh[name].fill.solid()
    sh[name].fill.fore_color.rgb = RGBColor(0xFB, 0xD9, 0xB5)
tick(s4, sh["Shape 11"])
fill(sh["Text 12"], [[("Working Prototype", True)]], color=ORANGE, size=10, align=PP_ALIGN.CENTER)
textbox(s4, 5.3, 1.28, 4.2, 0.28, [[("Live prototype: ", True, NAVY), (DEMO_LINK, True, ORANGE)]], size=10.5,
        align=PP_ALIGN.RIGHT, name="Demo link")
fill(sh["Text 19"], [
    [("Google Earth Engine", True), (" + web map app", False)],
    [("Data: ", True), ("CHIRPS, MODIS, Sentinel-1/2, WorldCover", False)],
    [("Next: drone survey", True), (" of flagged villages", False)],
], size=9, bullets=True)
fill(sh["Text 22"], [
    [("Data: ₹0", True), (" (free NASA, ESA and CHC data)", False)],
    [("Server: ~₹5,000 / month", True), (" for a whole state", False)],
    [("About ₹1-2 per village", True), (" per year", False)],
    [("District pilot: under ₹1 lakh", True), (", mostly field checks", False)],
], size=9, bullets=True)
fill(sh["Text 25"], [
    [("3 months: ", True), ("district pilot on the 2026 drought vs panchanamas", False)],
    [("6 months: ", True), ("drone surveys of flagged villages; Marathi app", False)],
    [("12 months: ", True), ("statewide, with land-record farm boundaries", False)],
], size=9, bullets=True)
# not yet engaged: named as the bodies we will approach for the pilot
fill(sh["Text 27"], [[("Partners to approach", True)]], color=NAVY, size=12 - BUMP)  # match the other headings
fill(sh["Text 28"], [
    "Relief & Rehabilitation Department, GoM",
    "Agriculture Department (taluka agriculture officers)",
    "MRSAC, Nagpur (state remote sensing centre)",
    "District collectors",
    "IIIT Nagpur",
], size=9, bullets=True)

# ---------------------------------------------------------------- slide 5: impact
s5, sh = prs.slides[4], S[5]
fill(sh["Text 6"], [[("48,926", True)]], color=NAVY, size=28, align=PP_ALIGN.CENTER)
fill(sh["Text 7"], [[("Villages covered", True)]], color=INK, size=11, align=PP_ALIGN.CENTER)
fill(sh["Text 8"], ["every village of Maharashtra, every season"], color=MUTED, size=9, align=PP_ALIGN.CENTER)
fill(sh["Text 10"], [[("Weeks → mins", True)]], color=NAVY, size=26, align=PP_ALIGN.CENTER)
fill(sh["Text 11"], [[("Saving in time", True)]], color=INK, size=11, align=PP_ALIGN.CENTER)
fill(sh["Text 12"], ["first screening vs a 1.5-2 month field survey"], color=MUTED, size=9, align=PP_ALIGN.CENTER)
fill(sh["Text 14"], [[("33 lakh+", True)]], color=NAVY, size=28, align=PP_ALIGN.CENTER)
fill(sh["Text 15"], [[("Farmer livelihoods supported", True)]], color=INK, size=11, align=PP_ALIGN.CENTER)
fill(sh["Text 16"], ["farmers in one 2025 relief round alone"], color=MUTED, size=9, align=PP_ALIGN.CENTER)
fill(sh["Text 19"], [
    "Relief and re-sowing come sooner for small farmers",
    "Fewer farms missed; officials verify the worst first",
    "Source-cited evidence reduces disputes",
], size=9.5, bullets=True)
fill(sh["Text 22"], [
    "Runs on free satellite data",
    "Low server cost fits a department budget",
    "Can be hosted by MRSAC; open-source code",
], size=9.5, bullets=True)
fill(sh["Text 25"], [
    "Any state: the drought manual is national and the data global",
    "Only needs that state's village boundaries",
    "Next: drone surveys for plot-level detail; rabi drought, hail",
], size=9.5, bullets=True)
textbox(s5, 0.5, 5.12, 5.4, 0.3, ["Sources: Govt of Maharashtra relief aid of ₹3,258 cr for 33.65 lakh farmers, Oct 2025 "
                                  "(The Hitavada); village count from Census 2001 boundaries (DataMeet). Server cost is "
                                  "an estimate."], size=6.5, color=MUTED)

# ---------------------------------------------------------------- slide 6: team
s6, sh = prs.slides[5], S[6]
team = [
    ("Text 6", "Text 7", "Text 8", "Text 9", "SC", "Sujal Chondekar", "Developer",
     "B.Tech CSE, IIIT Nagpur · satellite analysis, web app"),
    ("Text 12", "Text 13", "Text 14", "Text 15", "SH", "Dr. Swati Hira", "Faculty mentor",
     "Assistant Professor, CSE, IIIT Nagpur"),
]
for photo, name, role, qual, initials, n, r, q in team:
    fill(sh[photo], [[(initials, True)]], color=WHITE, size=11, align=PP_ALIGN.CENTER)
    fill(sh[name], [[(n, True)]], color=WHITE, size=10.5)
    fill(sh[role], [r], color=ORANGE, size=9)
    fill(sh[qual], [q], color=RGBColor(0xC9, 0xD2, 0xEC), size=8)
remove(s6, {"Shape 16", "Shape 17", "Text 18", "Text 19", "Text 20", "Text 21",
            "Shape 22", "Shape 23", "Text 24", "Text 25", "Text 26", "Text 27"})
sh["Text 28"].top = Inches(2.8)
fill(sh["Text 28"], ["Indian Institute of Information Technology, Nagpur"], color=RGBColor(0xC9, 0xD2, 0xEC), size=9.5)
for box in ("Shape 31", "Shape 33", "Shape 35"):
    tick(s6, sh[box], size=8)
fill(sh["Text 41"], ["A district pilot with the Relief & Rehabilitation and Agriculture Departments would let us "
                     "test the village list against panchanamas and make it the first step of every drought and "
                     "crop-loss assessment."], size=9)
sh["Text 42"].top, sh["Text 42"].height = Inches(3.25), Inches(0.6)
fill(sh["Text 42"], [[("Thank you  |  sujalchondhekar@gmail.com", True)],
                     [("Live demo: ", True), (DEMO_LINK, True)]], color=ORANGE, size=10.5, space_after=2)

prs.save(DECK)
print("saved", DECK)
